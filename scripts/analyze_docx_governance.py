#!/usr/bin/env python3
"""Analyze Chinese and English DOCX files for data/digital governance research.

The script reads DOCX files, tokenizes Chinese and English text, removes stopwords,
computes comparable word frequencies, and creates a small per-document LDA-style
topic summary by splitting each document into paragraph chunks.
"""

from __future__ import annotations

import argparse
import csv
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import jieba
import numpy as np
import pandas as pd
from docx import Document
from sklearn.decomposition import LatentDirichletAllocation
from sklearn.feature_extraction.text import CountVectorizer

ENGLISH_STOPWORDS = {
    "a", "about", "above", "according", "across", "after", "again", "against",
    "all", "also", "am", "among", "an", "and", "any", "are", "as", "at", "be",
    "because", "been", "before", "being", "below", "between", "both", "but", "by",
    "can", "could", "did", "do", "does", "doing", "down", "during", "each", "few",
    "for", "from", "further", "had", "has", "have", "having", "he", "her", "here",
    "hers", "herself", "him", "himself", "his", "how", "i", "if", "in", "into", "is",
    "it", "its", "itself", "just", "may", "more", "most", "must", "my", "myself", "no",
    "nor", "not", "now", "of", "off", "on", "once", "only", "or", "other", "our",
    "ours", "ourselves", "out", "over", "own", "same", "shall", "she", "should", "so",
    "some", "such", "than", "that", "the", "their", "theirs", "them", "themselves",
    "then", "there", "these", "they", "this", "those", "through", "to", "too", "under",
    "until", "up", "very", "was", "we", "were", "what", "when", "where", "which", "while",
    "who", "whom", "why", "will", "with", "within", "would", "you", "your", "yours",
    "yourself", "yourselves",
}

CHINESE_STOPWORDS = {
    "的", "了", "和", "与", "及", "或", "并", "在", "对", "为", "是", "于", "中", "由", "将",
    "应", "该", "其", "本", "等", "以", "但", "而", "且", "就", "都", "也", "被", "把", "从",
    "到", "向", "根据", "有关", "相关", "进行", "可以", "不得", "以及", "或者", "如果", "需要",
    "通过", "采取", "包括", "具有", "关于", "一个", "一种", "第", "条", "款", "项", "章",
}

TOKEN_RE = re.compile(r"[a-zA-Z][a-zA-Z\-']*[a-zA-Z]")
CJK_RE = re.compile(r"[\u4e00-\u9fff]")


@dataclass(frozen=True)
class DocRecord:
    name: str
    path: Path
    language: str
    paragraphs: list[str]
    tokens: list[str]


def read_docx(path: Path) -> list[str]:
    document = Document(path)
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def guess_language(text: str) -> str:
    cjk_count = len(CJK_RE.findall(text))
    latin_count = len(re.findall(r"[A-Za-z]", text))
    return "zh" if cjk_count > latin_count else "en"


def parse_language_map(raw_map: str | None) -> dict[str, str]:
    if not raw_map:
        return {}
    language_map: dict[str, str] = {}
    for item in raw_map.split(","):
        if not item.strip():
            continue
        key, separator, value = item.partition(":")
        if not separator or value not in {"zh", "en"}:
            raise ValueError("Language map must look like 'china:zh,eu:en'.")
        language_map[key.strip().lower()] = value
    return language_map


def load_extra_stopwords(path: Path | None) -> set[str]:
    if path is None or not path.exists():
        return set()
    words: set[str] = set()
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            word = line.strip()
            if word and not word.startswith("#"):
                words.add(word.lower())
    return words


def tokenize_english(text: str, stopwords: set[str], min_token_len: int) -> list[str]:
    tokens: list[str] = []
    for match in TOKEN_RE.finditer(text.lower()):
        token = match.group().strip("-'")
        if len(token) >= min_token_len and token not in stopwords:
            tokens.append(token)
    return tokens


def tokenize_chinese(text: str, stopwords: set[str], min_token_len: int) -> list[str]:
    tokens: list[str] = []
    for raw_token in jieba.cut(text):
        token = raw_token.strip().lower()
        if not token or token in stopwords:
            continue
        if re.fullmatch(r"[\W_]+", token):
            continue
        if CJK_RE.search(token):
            if len(token) >= min_token_len:
                tokens.append(token)
        elif len(token) >= min_token_len and token not in ENGLISH_STOPWORDS:
            tokens.append(token)
    return tokens


def tokenize(text: str, language: str, extra_stopwords: set[str], min_token_len: int) -> list[str]:
    if language == "zh":
        return tokenize_chinese(text, CHINESE_STOPWORDS | extra_stopwords, min_token_len)
    return tokenize_english(text, ENGLISH_STOPWORDS | extra_stopwords, min_token_len)


def load_documents(
    input_dir: Path,
    language_map: dict[str, str],
    extra_stopwords: set[str],
    min_token_len: int,
) -> list[DocRecord]:
    records: list[DocRecord] = []
    for path in sorted(input_dir.glob("*.docx")):
        paragraphs = read_docx(path)
        text = "\n".join(paragraphs)
        stem = path.stem.lower()
        language = language_map.get(stem, guess_language(text))
        tokens = tokenize(text, language, extra_stopwords, min_token_len)
        records.append(DocRecord(path.stem, path, language, paragraphs, tokens))
    if not records:
        raise FileNotFoundError(f"No .docx files found in {input_dir}")
    return records


def write_frequency_outputs(records: list[DocRecord], output_dir: Path, top_n: int) -> None:
    rows: list[dict[str, object]] = []
    summary_rows: list[dict[str, object]] = []
    for record in records:
        counts = Counter(record.tokens)
        total_tokens = sum(counts.values())
        summary_rows.append(
            {
                "document": record.name,
                "language": record.language,
                "source_file": str(record.path),
                "total_tokens_after_stopword_removal": total_tokens,
                "unique_tokens": len(counts),
            }
        )
        for token, count in counts.items():
            rows.append(
                {
                    "document": record.name,
                    "language": record.language,
                    "token": token,
                    "count": count,
                    "freq_per_10k": (count / total_tokens * 10000) if total_tokens else 0,
                }
            )

    frequencies = pd.DataFrame(rows).sort_values(
        ["document", "freq_per_10k", "count", "token"], ascending=[True, False, False, True]
    )
    frequencies.to_csv(output_dir / "word_frequencies.csv", index=False, quoting=csv.QUOTE_MINIMAL)

    top_words = frequencies.groupby("document", group_keys=False).head(top_n)
    top_words.to_csv(output_dir / "top_words_by_document.csv", index=False, quoting=csv.QUOTE_MINIMAL)

    comparison = top_words.pivot_table(
        index="document", columns="token", values="freq_per_10k", fill_value=0, aggfunc="sum"
    )
    comparison.to_csv(output_dir / "comparison_matrix.csv")

    pd.DataFrame(summary_rows).to_csv(output_dir / "document_summary.csv", index=False)


def paragraph_chunks(record: DocRecord, extra_stopwords: set[str], min_token_len: int, min_topic_tokens: int) -> list[str]:
    chunks: list[str] = []
    current_tokens: list[str] = []
    for paragraph in record.paragraphs:
        paragraph_tokens = tokenize(paragraph, record.language, extra_stopwords, min_token_len)
        if not paragraph_tokens:
            continue
        current_tokens.extend(paragraph_tokens)
        if len(current_tokens) >= min_topic_tokens:
            chunks.append(" ".join(current_tokens))
            current_tokens = []
    if current_tokens:
        if chunks:
            chunks[-1] = f"{chunks[-1]} {' '.join(current_tokens)}"
        elif len(current_tokens) >= max(5, min_topic_tokens // 2):
            chunks.append(" ".join(current_tokens))
    return chunks


def identity_analyzer(text: str) -> list[str]:
    return text.split()


def write_lda_outputs(
    records: list[DocRecord],
    output_dir: Path,
    topics: int,
    words_per_topic: int,
    extra_stopwords: set[str],
    min_token_len: int,
    min_topic_tokens: int,
) -> None:
    topic_rows: list[dict[str, object]] = []
    assignment_rows: list[dict[str, object]] = []

    for record in records:
        chunks = paragraph_chunks(record, extra_stopwords, min_token_len, min_topic_tokens)
        if len(chunks) < 2:
            topic_rows.append(
                {
                    "document": record.name,
                    "topic": "skipped",
                    "top_terms": "",
                    "note": "Not enough paragraph chunks for LDA. Try lowering --min-topic-tokens.",
                }
            )
            continue

        vectorizer = CountVectorizer(analyzer=identity_analyzer, min_df=1)
        matrix = vectorizer.fit_transform(chunks)
        feature_names = np.array(vectorizer.get_feature_names_out())
        n_topics = min(topics, max(1, len(chunks) - 1), matrix.shape[1])
        lda = LatentDirichletAllocation(n_components=n_topics, random_state=42, learning_method="batch")
        distributions = lda.fit_transform(matrix)

        for topic_index, topic_weights in enumerate(lda.components_, start=1):
            top_indices = topic_weights.argsort()[::-1][:words_per_topic]
            topic_rows.append(
                {
                    "document": record.name,
                    "topic": topic_index,
                    "top_terms": ", ".join(feature_names[top_indices]),
                    "note": "",
                }
            )

        for chunk_index, topic_distribution in enumerate(distributions, start=1):
            best_topic = int(topic_distribution.argmax() + 1)
            assignment_rows.append(
                {
                    "document": record.name,
                    "chunk": chunk_index,
                    "dominant_topic": best_topic,
                    "dominant_topic_probability": float(topic_distribution.max()),
                    "token_count": len(chunks[chunk_index - 1].split()),
                    "chunk_preview": " ".join(chunks[chunk_index - 1].split()[:40]),
                }
            )

    pd.DataFrame(topic_rows).to_csv(output_dir / "lda_topics.csv", index=False)
    pd.DataFrame(assignment_rows).to_csv(output_dir / "lda_chunk_topics.csv", index=False)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Tokenize DOCX governance texts and compare word frequencies/topics.")
    parser.add_argument("--input-dir", type=Path, default=Path("data"), help="Folder containing .docx files.")
    parser.add_argument("--output-dir", type=Path, default=Path("outputs"), help="Folder for CSV outputs.")
    parser.add_argument("--language-map", help="Optional map such as 'china:zh,eu:en,usa:en,un:en'.")
    parser.add_argument("--extra-stopwords", type=Path, default=Path("stopwords/custom_stopwords.txt"))
    parser.add_argument("--top-n", type=int, default=50, help="Number of top words exported per document.")
    parser.add_argument("--topics", type=int, default=3, help="Maximum number of LDA topics per document.")
    parser.add_argument("--words-per-topic", type=int, default=12, help="Number of keywords shown for each topic.")
    parser.add_argument("--min-token-len", type=int, default=2, help="Minimum token length after segmentation.")
    parser.add_argument("--min-topic-tokens", type=int, default=30, help="Minimum tokens per paragraph chunk for LDA.")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    language_map = parse_language_map(args.language_map)
    extra_stopwords = load_extra_stopwords(args.extra_stopwords)
    records = load_documents(args.input_dir, language_map, extra_stopwords, args.min_token_len)
    write_frequency_outputs(records, args.output_dir, args.top_n)
    write_lda_outputs(
        records,
        args.output_dir,
        args.topics,
        args.words_per_topic,
        extra_stopwords,
        args.min_token_len,
        args.min_topic_tokens,
    )
    print(f"Analyzed {len(records)} DOCX file(s). Outputs written to {args.output_dir}.")


if __name__ == "__main__":
    main()
